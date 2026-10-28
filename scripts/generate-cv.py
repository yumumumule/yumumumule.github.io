"""Create an editable public CV from curated content; validate an exported PDF.

Requires python-docx, pypdf and reportlab. PDF rendering is a separate visual QA step.
Use the Codex bundled Python runtime when running inside Codex.
"""
import argparse
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'content/cv.json').read_text(encoding='utf-8'))

def hyperlink(paragraph, label, url, *, bold=False, size=None):
    relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    element = OxmlElement('w:hyperlink')
    element.set(qn('r:id'), relationship)
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '874330'); props.append(color)
    if bold: props.append(OxmlElement('w:b'))
    if size:
        font_size = OxmlElement('w:sz'); font_size.set(qn('w:val'), str(round(size*2))); props.append(font_size)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    element.append(run); paragraph._p.append(element)

def generate_docx(destination):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210); section.page_height = Mm(297)
    section.top_margin = Inches(.48); section.bottom_margin = Inches(.48)
    section.left_margin = Inches(.59); section.right_margin = Inches(.59)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.02
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.widow_control = True
    for name in ['Title','Subtitle','Heading 1','Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0,0,0)
        style.font.underline = False
    title = doc.styles['Title']
    title.font.name = 'Georgia'; title.font.size = Pt(27); title.font.bold = False
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.line_spacing = 1
    heading = doc.styles['Heading 1']
    heading.font.size = Pt(11); heading.font.bold = True
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.line_spacing = 1
    width = section.page_width - section.left_margin - section.right_margin

    def p(text='', *, size=None, italic=False, after=2, keep=False):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.keep_with_next = keep
        if text:
            run=paragraph.add_run(text); run.italic=italic
            if size: run.font.size=Pt(size)
        return paragraph

    def row(left, right):
        paragraph=p(keep=True)
        paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run(left).bold=True
        paragraph.add_run('\t'+right).font.size=Pt(9.5)
        return paragraph

    doc.add_paragraph(DATA['name'], 'Title')
    p('CUHK AIST undergraduate | AI/ML and software engineering internships - Summer 2027',size=9.5,after=3)
    contacts=p(after=4)
    for index,(label,url) in enumerate([(DATA['email'],'mailto:'+DATA['email']),('Website',DATA['website']),('GitHub',DATA['github']),('OpenReview',DATA['openreview'])]):
        if index: contacts.add_run('   |   ').font.size=Pt(9.5)
        hyperlink(contacts,label,url,size=9.5)
    p(DATA['interests'],size=9.5,after=1)

    doc.add_heading('Education',1)
    education=DATA['education']
    row(education['institution'],education['period'])
    p(education['degree'])
    p(education['record'],size=10)
    p('Relevant coursework: '+education['coursework'],size=9.5)

    doc.add_heading('Research Experience',1)
    research=DATA['research']
    row(research['lab'],research['period'])
    p(research['role'],size=10,italic=True)
    for detail in research['details']: p(detail,size=10)

    doc.add_heading('Papers and Preprints',1)
    for index,paper in enumerate(DATA['papers']):
        paragraph=p(after=1,keep=True)
        hyperlink(paragraph,paper['title'],paper['url'],bold=True,size=10.5)
        # Name emphasis preserves the published author order; no special author role is implied.
        authors=p(after=0,keep=True)
        for part_index,part in enumerate(paper['authors'].split('Hanlin Zhao')):
            if part_index:
                run=authors.add_run('Hanlin Zhao'); run.bold=True; run.font.size=Pt(9.3)
            run=authors.add_run(part); run.font.size=Pt(9.3)
        p(paper['status'],size=9.3,after=4 if index<len(DATA['papers'])-1 else 1)

    doc.add_heading('Selected Engineering Projects',1)
    for project in DATA['projects']:
        paragraph=p(after=1,keep=True)
        hyperlink(paragraph,project['name'],project['url'],bold=True,size=10.5)
        paragraph.add_run(' | '+project['tech']).font.size=Pt(9.3)
        p(project['description'],size=10,after=4)

    doc.add_heading('Selected Honors and Activities',1)
    for honor in DATA['honors']: p(honor,size=10,after=1)
    activities=p(after=1)
    for index,activity in enumerate(DATA['activities']):
        if index: activities.add_run(' | ')
        label=activity['role']
        if label=='Shaw College Student Ambassador': label='Shaw College Ambassador'
        if label=='Chinese Debate Team Member': label='Chinese Debate Team'
        run=activities.add_run(label);run.font.size=Pt(9.5)

    doc.add_heading('Technical Skills',1)
    p(DATA['skills'],size=10,after=0)
    doc.core_properties.title='Hanlin Zhao Curriculum Vitae'
    doc.core_properties.author='Hanlin Zhao'
    doc.core_properties.last_modified_by=''
    doc.core_properties.subject=''
    doc.core_properties.comments=''
    destination.parent.mkdir(parents=True,exist_ok=True)
    doc.save(destination)
    print('Created editable public CV:',destination)

def finish_pdf(source,destination,docx_source):
    from pypdf import PdfReader,PdfWriter
    reader=PdfReader(source)
    if len(reader.pages)!=1: raise ValueError(f'Expected a one-page CV, got {len(reader.pages)} pages')
    writer=PdfWriter()
    writer.clone_document_from_reader(reader)
    writer._info.get_object().clear()
    writer.add_metadata({'/Title':'Hanlin Zhao Curriculum Vitae','/Author':'Hanlin Zhao'})
    destination.parent.mkdir(parents=True,exist_ok=True)
    with destination.open('wb') as file: writer.write(file)
    final=PdfReader(destination)
    text='\n'.join(page.extract_text() for page in final.pages)
    def normalize(value):
        for ligature,letters in [('ﬁ','fi'),('ﬂ','fl'),('ﬀ','ff'),('ﬃ','ffi'),('ﬄ','ffl')]: value=value.replace(ligature,letters)
        return re.sub(r'\s+','',value)
    expected=Document(docx_source)
    for para in expected.element.body.iter(qn('w:p')):
        value=''.join(node.text or '' for node in para.iter(qn('w:t')))
        if value.strip() and normalize(value) not in normalize(text): raise ValueError('Missing PDF text: '+value)
    assert DATA['email'] in text
    assert 'link.cuhk.edu.hk' not in text
    assert not re.search(r'\+\s*(?:852|86)\s*\d',text)
    assert not re.search(r'\b\d{10}@',text)
    assert 'uARPRKfrqd' not in text
    urls=[]
    for page in final.pages:
        for ref in page.get('/Annots',[]):
            action=ref.get_object().get('/A')
            if action and action.get('/URI'): urls.append(str(action['/URI']))
    for paper in DATA['papers']: assert paper['url'] in urls, 'Missing PDF paper hyperlink'
    (source.parent/'final-text.txt').write_text(text,encoding='utf-8')
    print(f'Validated PDF: one page, {destination.stat().st_size} bytes, source text retained, paper hyperlinks intact, privacy checks passed.')

def render_pdf(destination):
    """Lay out the same curated content without relying on desktop Office."""
    from html import escape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Flowable
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    fonts=Path('C:/Windows/Fonts')
    for name,filename in [('CV','calibri.ttf'),('CV-Bold','calibrib.ttf'),('CV-Italic','calibrii.ttf'),('CV-BoldItalic','calibriz.ttf'),('CV-Title','georgia.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(fonts/filename)))
    pdfmetrics.registerFontFamily('CV',normal='CV',bold='CV-Bold',italic='CV-Italic',boldItalic='CV-BoldItalic')
    styles={
        'title':ParagraphStyle('title',fontName='CV-Title',fontSize=29,leading=33,spaceAfter=5,textColor='#000000'),
        'header':ParagraphStyle('header',fontName='CV',fontSize=10,leading=12.8,spaceAfter=3,textColor='#000000'),
        'heading':ParagraphStyle('heading',fontName='CV-Bold',fontSize=11.5,leading=14,spaceBefore=10,spaceAfter=4,keepWithNext=True,textColor='#000000'),
        'body':ParagraphStyle('body',fontName='CV',fontSize=11,leading=13.3,spaceAfter=2,textColor='#000000'),
        'small':ParagraphStyle('small',fontName='CV',fontSize=9.5,leading=11.8,spaceAfter=1,textColor='#000000'),
        'paper':ParagraphStyle('paper',fontName='CV-Bold',fontSize=11,leading=13.2,spaceAfter=1,keepWithNext=True,textColor='#000000')
    }
    def link(label,url): return f'<link href="{escape(url,quote=True)}" color="#874330">{escape(label)}</link>'
    story=[]
    def para(text,kind='body',markup=False,after=None,keep=False):
        style=ParagraphStyle('item',parent=styles[kind])
        if after is not None: style.spaceAfter=after
        if keep: style.keepWithNext=True
        story.append(Paragraph(text if markup else escape(text),style))
    class Row(Flowable):
        def __init__(self,left,right):
            super().__init__();self.left=left;self.right=right;self.height=14.5;self.keepWithNext=True
        def wrap(self,availableWidth,availableHeight):self.width=availableWidth;return self.width,self.height
        def draw(self):
            self.canv.setFillColorRGB(0,0,0);self.canv.setFont('CV-Bold',11)
            self.canv.drawString(0,2,self.left);self.canv.setFont('CV',9.5)
            self.canv.drawRightString(self.width,2,self.right)
    para(DATA['name'],'title')
    para('CUHK AIST undergraduate | AI/ML and software engineering internships - Summer 2027','header')
    para('   |   '.join([link(DATA['email'],'mailto:'+DATA['email']),link('Website',DATA['website']),link('GitHub',DATA['github']),link('OpenReview',DATA['openreview'])]),'header',True)
    para(DATA['interests'],'header',after=1)
    para('Education','heading')
    education=DATA['education'];story.append(Row(education['institution'],education['period']))
    para(education['degree']);para(education['record']);para('Relevant coursework: '+education['coursework'],'small')
    para('Research Experience','heading')
    research=DATA['research'];story.append(Row(research['lab'],research['period']))
    para('<i>'+escape(research['role'])+'</i>','body',True)
    for detail in research['details']:para(detail)
    para('Papers and Preprints','heading')
    for paper in DATA['papers']:
        para(link(paper['title'],paper['url']),'paper',True)
        para(escape(paper['authors']).replace('Hanlin Zhao','<b>Hanlin Zhao</b>'),'small',True,after=0,keep=True)
        para(paper['status'],'small',after=4)
    para('Selected Engineering Projects','heading')
    for project in DATA['projects']:
        para(link(project['name'],project['url'])+' <font name="CV" size="9.3">| '+escape(project['tech'])+'</font>','paper',True)
        para(project['description'],'body',after=4)
    para('Selected Honors and Activities','heading')
    for honor in DATA['honors']:para(honor,'body',after=1)
    para('Shaw College Ambassador | ICP Mentor | Chinese Debate Team','small',after=1)
    para('Technical Skills','heading');para(DATA['skills'],'body',after=0)
    raw=ROOT/'artifacts/cv-refresh/Hanlin_Zhao_CV_Layout.pdf'
    raw.parent.mkdir(parents=True,exist_ok=True)
    layout=SimpleDocTemplate(str(raw),pagesize=A4,leftMargin=42.5,rightMargin=42.5,topMargin=34.5,bottomMargin=34.5,title='Hanlin Zhao Curriculum Vitae',author='Hanlin Zhao')
    layout.build(story)
    source_docx=ROOT/'artifacts/cv-refresh/Hanlin_Zhao_CV_Updated.docx'
    finish_pdf(raw,destination,source_docx)

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--docx-out',type=Path)
    parser.add_argument('--pdf-in',type=Path)
    parser.add_argument('--pdf-out',type=Path,default=ROOT/'assets/Hanlin_Zhao_CV.pdf')
    parser.add_argument('--docx-source',type=Path)
    parser.add_argument('--render-pdf',action='store_true')
    args=parser.parse_args()
    if args.render_pdf: render_pdf(args.pdf_out)
    elif args.docx_out: generate_docx(args.docx_out)
    elif args.pdf_in and args.docx_source: finish_pdf(args.pdf_in,args.pdf_out,args.docx_source)
    else: parser.error('Provide --docx-out, or --pdf-in and --docx-source')
